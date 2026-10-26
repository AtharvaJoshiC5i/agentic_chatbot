"""Renderers for the Phase 4 unstructured corpus (PDF, DOCX, EML, TXT).

Documents are described as a small block model so the same content can be
written to PDF (reportlab) or DOCX (python-docx) with a real text layer.
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timedelta, timezone
from email.message import EmailMessage
from email.utils import format_datetime
from pathlib import Path
from xml.sax.saxutils import escape

IST = timezone(timedelta(hours=5, minutes=30))
FIXED_CREATED = datetime(2026, 3, 20, 12, 0, 0)
AUTHOR = "ShopSphere Retail"

# Block kinds: ("h", text) | ("p", text) | ("bullets", [text]) |
# ("kv", [(key, value)]) | ("table", [header], [[cell]])
Block = tuple


@dataclass
class Doc:
    title: str
    subtitle: str = ""
    header: list[tuple[str, str]] = field(default_factory=list)
    blocks: list[Block] = field(default_factory=list)
    footer: str = "Internal use only — ShopSphere Retail."

    def h(self, text: str) -> "Doc":
        self.blocks.append(("h", text))
        return self

    def p(self, text: str) -> "Doc":
        self.blocks.append(("p", text))
        return self

    def bullets(self, items: list[str]) -> "Doc":
        self.blocks.append(("bullets", list(items)))
        return self

    def kv(self, pairs: list[tuple[str, str]]) -> "Doc":
        self.blocks.append(("kv", list(pairs)))
        return self

    def table(self, header: list[str], rows: list[list[str]]) -> "Doc":
        self.blocks.append(("table", list(header), [list(r) for r in rows]))
        return self

    def plain_text(self) -> str:
        """Flatten to text (used for word counts and content checks)."""
        parts = [self.title, self.subtitle]
        parts += [f"{k}: {v}" for k, v in self.header]
        for block in self.blocks:
            kind = block[0]
            if kind in ("h", "p"):
                parts.append(block[1])
            elif kind == "bullets":
                parts += block[1]
            elif kind == "kv":
                parts += [f"{k}: {v}" for k, v in block[1]]
            elif kind == "table":
                parts.append(" | ".join(block[1]))
                parts += [" | ".join(r) for r in block[2]]
        parts.append(self.footer)
        return "\n".join(p for p in parts if p)


# --------------------------------------------------------------------------- PDF

# Standard Type 1 fonts carry no ToUnicode map, so text extractors (pypdf) cannot
# decode non-ASCII glyphs. PDF text is normalised to ASCII to keep it extractable.
_PDF_REPLACEMENTS = {
    "₹": "INR ",
    "≥": ">=",
    "≤": "<=",
    "→": "->",
    "×": "x",
    "—": "-",
    "–": "-",
    "−": "-",
    "·": "|",
    "“": '"',
    "”": '"',
    "‘": "'",
    "’": "'",
    " ": " ",
}


def _pdf_safe(text: str) -> str:
    for src, dst in _PDF_REPLACEMENTS.items():
        text = text.replace(src, dst)
    if not text.isascii():
        bad = sorted({ch for ch in text if not ch.isascii()})
        raise ValueError(f"Add a PDF replacement for non-ASCII characters {bad!r} in: {text[:80]!r}")
    return text


def write_pdf(doc: Doc, path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        ListFlowable,
        ListItem,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=9.5, leading=13)
    cell = ParagraphStyle("cell", parent=body, fontSize=8.5, leading=11)
    cell_bold = ParagraphStyle("cellb", parent=cell, fontName="Helvetica-Bold")
    title = ParagraphStyle("t", parent=styles["Title"], fontSize=16, leading=20, spaceAfter=4)
    subtitle = ParagraphStyle("st", parent=body, textColor=colors.HexColor("#444444"), spaceAfter=8)
    heading = ParagraphStyle("hd", parent=styles["Heading2"], fontSize=11.5, leading=14, spaceBefore=8)

    def para(text: str, style=body) -> Paragraph:
        return Paragraph(escape(_pdf_safe(text)), style)

    def kv_table(pairs: list[tuple[str, str]]) -> Table:
        data = [[para(k, cell_bold), para(v, cell)] for k, v in pairs]
        t = Table(data, colWidths=[52 * mm, 118 * mm])
        t.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
                    ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EEF1F5")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
            )
        )
        return t

    story = [para(doc.title, title)]
    if doc.subtitle:
        story.append(para(doc.subtitle, subtitle))
    if doc.header:
        story += [kv_table(doc.header), Spacer(1, 6)]
    for block in doc.blocks:
        kind = block[0]
        if kind == "h":
            story.append(para(block[1], heading))
        elif kind == "p":
            story.append(para(block[1]))
            story.append(Spacer(1, 3))
        elif kind == "bullets":
            story.append(
                ListFlowable(
                    [ListItem(para(item), leftIndent=10) for item in block[1]],
                    bulletType="bullet",
                    start="-",
                    leftIndent=12,
                )
            )
            story.append(Spacer(1, 3))
        elif kind == "kv":
            story += [kv_table(block[1]), Spacer(1, 4)]
        elif kind == "table":
            header, rows = block[1], block[2]
            data = [[para(h, cell_bold) for h in header]] + [[para(c, cell) for c in r] for r in rows]
            width = 170 * mm / len(header)
            t = Table(data, colWidths=[width] * len(header), repeatRows=1)
            t.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDE3EA")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ]
                )
            )
            story += [t, Spacer(1, 4)]

    footer_text = _pdf_safe(doc.footer)

    def on_page(canvas, document) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(20 * mm, 10 * mm, footer_text)
        canvas.drawRightString(190 * mm, 10 * mm, f"Page {document.page}")
        canvas.restoreState()

    path.parent.mkdir(parents=True, exist_ok=True)
    pdf = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=_pdf_safe(doc.title),
        author=AUTHOR,
        subject=_pdf_safe(doc.subtitle or doc.title),
        creator="ShopSphere synthetic document generator",
        invariant=1,
    )
    pdf.build(story, onFirstPage=on_page, onLaterPages=on_page)


# -------------------------------------------------------------------------- DOCX


def write_docx(doc: Doc, path: Path) -> None:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    props = document.core_properties
    props.title = doc.title
    props.author = AUTHOR
    props.subject = doc.subtitle or doc.title
    props.created = FIXED_CREATED
    props.modified = FIXED_CREATED
    props.last_modified_by = AUTHOR

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    def add_kv(pairs: list[tuple[str, str]]) -> None:
        table = document.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for key, value in pairs:
            row = table.add_row().cells
            row[0].text = key
            row[1].text = value
            for run in row[0].paragraphs[0].runs:
                run.bold = True
        document.add_paragraph()

    document.add_heading(doc.title, level=0)
    if doc.subtitle:
        document.add_paragraph(doc.subtitle).runs[0].italic = True
    if doc.header:
        add_kv(doc.header)
    for block in doc.blocks:
        kind = block[0]
        if kind == "h":
            document.add_heading(block[1], level=2)
        elif kind == "p":
            document.add_paragraph(block[1])
        elif kind == "bullets":
            for item in block[1]:
                document.add_paragraph(item, style="List Bullet")
        elif kind == "kv":
            add_kv(block[1])
        elif kind == "table":
            header, rows = block[1], block[2]
            table = document.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            for i, text in enumerate(header):
                table.rows[0].cells[i].text = text
                for run in table.rows[0].cells[i].paragraphs[0].runs:
                    run.bold = True
            for r in rows:
                cells = table.add_row().cells
                for i, text in enumerate(r):
                    cells[i].text = text
            document.add_paragraph()

    footer = document.sections[0].footer.paragraphs[0]
    footer.text = doc.footer

    # Re-pack with fixed zip timestamps so regenerated files are byte-identical.
    buffer = io.BytesIO()
    document.save(buffer)
    path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(buffer) as src, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            info = zipfile.ZipInfo(item.filename, date_time=FIXED_CREATED.timetuple()[:6])
            info.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(info, src.read(item.filename))


# ----------------------------------------------------------------------- EML/TXT


@dataclass
class Email:
    sender: str
    to: str
    subject: str
    sent_at: datetime
    body: str
    message_id: str
    cc: str = ""
    in_reply_to: str = ""
    references: list[str] = field(default_factory=list)
    extra_headers: dict[str, str] = field(default_factory=dict)


def write_eml(email: Email, path: Path) -> None:
    msg = EmailMessage()
    msg["From"] = email.sender
    msg["To"] = email.to
    if email.cc:
        msg["Cc"] = email.cc
    msg["Subject"] = email.subject
    msg["Date"] = format_datetime(email.sent_at.replace(tzinfo=IST))
    msg["Message-ID"] = email.message_id
    if email.in_reply_to:
        msg["In-Reply-To"] = email.in_reply_to
    if email.references:
        msg["References"] = " ".join(email.references)
    for key, value in email.extra_headers.items():
        msg[key] = value
    # 8bit keeps the body human- and parser-readable (no quoted-printable soft breaks).
    msg.set_content(email.body.strip() + "\n", charset="utf-8", cte="8bit")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(msg.as_bytes())


def write_txt(text: str, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text.strip() + "\n", encoding="utf-8", newline="\n")
