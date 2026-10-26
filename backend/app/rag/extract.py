"""Extract plain text from unstructured files."""

from __future__ import annotations

import email
import re
from dataclasses import dataclass
from email import policy as email_policy
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader


@dataclass(frozen=True)
class TextSegment:
    """A contiguous text block with optional page/section hint."""

    text: str
    page: int | None = None
    section: str | None = None


@dataclass(frozen=True)
class ExtractedDocument:
    segments: tuple[TextSegment, ...]

    @property
    def full_text(self) -> str:
        return "\n\n".join(segment.text for segment in self.segments if segment.text.strip())


def extract_document(path: Path) -> ExtractedDocument:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix == ".txt":
        return ExtractedDocument((TextSegment(path.read_text(encoding="utf-8")),))
    if suffix == ".eml":
        return _extract_eml(path)
    raise ValueError(f"Unsupported document format: {path.suffix}")


def _extract_pdf(path: Path) -> ExtractedDocument:
    reader = PdfReader(str(path))
    segments: list[TextSegment] = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        text = _normalize_whitespace(text)
        if text:
            segments.append(TextSegment(text=text, page=index, section=f"Page {index}"))
    if not segments:
        raise ValueError(f"No extractable text in PDF: {path}")
    return ExtractedDocument(tuple(segments))


def _extract_docx(path: Path) -> ExtractedDocument:
    document = DocxDocument(str(path))
    segments: list[TextSegment] = []
    for paragraph in document.paragraphs:
        text = _normalize_whitespace(paragraph.text)
        if not text:
            continue
        style = paragraph.style.name if paragraph.style else "Body"
        segments.append(TextSegment(text=text, section=style))
    if not segments:
        raise ValueError(f"No extractable text in DOCX: {path}")
    return ExtractedDocument(tuple(segments))


def _extract_eml(path: Path) -> ExtractedDocument:
    raw = path.read_bytes()
    message = email.message_from_bytes(raw, policy=email_policy.default)
    headers = [
        f"Subject: {message.get('Subject', '')}",
        f"From: {message.get('From', '')}",
        f"To: {message.get('To', '')}",
        f"Date: {message.get('Date', '')}",
    ]
    body = message.get_body(preferencelist=("plain", "html"))
    body_text = ""
    if body is not None:
        body_text = body.get_content()
        if body.get_content_type() == "text/html":
            body_text = re.sub(r"<[^>]+>", " ", body_text)
    combined = "\n".join(headers) + "\n\n" + _normalize_whitespace(body_text)
    return ExtractedDocument((TextSegment(text=combined, section="Email"),))


def _normalize_whitespace(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
